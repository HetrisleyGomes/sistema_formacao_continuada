from django.contrib.admin.views.decorators import staff_member_required
from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Sum
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from formacoes.models import Formacao
from turmas.models import Turma
from encontros.models import Encontro
from inscricoes.models import Inscricao
from inscricoes.services import calcular_frequencia

from datetime import datetime
from html import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)

@staff_member_required
def relatorio_participantes(request):
    formacao_id = request.GET.get("formacao", "")
    turma_id = request.GET.get("turma", "")

    formacoes = Formacao.objects.all()
    turmas = Turma.objects.all()

    inscricoes = Inscricao.objects.filter(
        situacao = Inscricao.Situacao.ATIVA
    ).select_related(
        "cursista",
        "turma",
        "turma__formacao",
    )

    if formacao_id.isdigit():
        inscricoes = inscricoes.filter(
            turma__formacao_id=int(formacao_id)
        )
        turmas = turmas.filter(
            formacao_id=int(formacao_id)
        )

    if turma_id.isdigit():
        inscricoes = inscricoes.filter(
            turma_id=int(turma_id)
        )

    resultados = []

    for inscricao in inscricoes:
        calculo = calcular_frequencia(inscricao)

        resultados.append({
            "inscricao": inscricao,
            **calculo,
        })

    concluintes = [
        item for item in resultados
        if item["concluinte"]
    ]

    contexto = {
        "formacoes": formacoes,
        "turmas": turmas,
        "resultados": resultados,
        "concluintes": concluintes,
        "total_participantes": len(resultados),
        "total_concluintes": len(concluintes),
        "formacao_selecionada": formacao_id,
        "turma_selecionada": turma_id,
    }

    return render(
        request,
        "relatorios/participantes.html",
        contexto,
    )

@staff_member_required
def export_participantes_excel(request):
    formacao_id = request.GET.get("formacao", "")
    turma_id = request.GET.get("turma", "")

    inscricoes = Inscricao.objects.filter(
        situacao=Inscricao.Situacao.ATIVA
    ).select_related(
        "cursista",
        "cursista__escola",
        "turma",
        "turma__formacao",
    )

    if formacao_id.isdigit():
        inscricoes = inscricoes.filter(
            turma__formacao_id=int(formacao_id)
        )

    if turma_id.isdigit():
        inscricoes = inscricoes.filter(
            turma_id=int(turma_id)
        )

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Participantes"

    cabecalhos = [
        "Nome completo",
        "CPF",
        "Escola",
        "Função",
        "Formação",
        "Turma",
        "Horas realizadas",
        "Horas frequentadas",
        "Frequência (%)",
        "Resultado",
    ]

    planilha.append(cabecalhos)

    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(
            fill_type="solid", fgColor="176B39"
        )
        celula.alignment = Alignment(
            horizontal="center"
        )

    for inscricao in inscricoes:
        calculo = calcular_frequencia(inscricao)
        cursista = inscricao.cursista

        if calculo["horas_realizadas"] == 0:
            resultado = "Sem encontros realizados"
        elif calculo["concluinte"]:
            resultado = "Concluinte"
        else:
            resultado = "Não concluinte"

        planilha.append([
            cursista.nome_completo,
            cursista.cpf,
            cursista.escola.nome,
            cursista.get_funcao_display(),
            inscricao.turma.formacao.titulo,
            inscricao.turma.nome,
            float(calculo["horas_realizadas"]),
            float(calculo["horas_frequentadas"]),
            float(calculo["percentual_frequencia"]),
            resultado,
        ])

    larguras = {
        "A": 32,
        "B": 16,
        "C": 30,
        "D": 24,
        "E": 36,
        "F": 24,
        "G": 18,
        "H": 21,
        "I": 18,
        "J": 28,
    }

    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions

    resposta = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    resposta["Contente-Disposition"] = (
        'attachment; filename="relatorio_participantes.xlsx'
    )

    workbook.save(resposta)
    return resposta

@staff_member_required
def export_participantes_pdf(request):
    formacao_id = request.GET.get("formacao", "")
    turma_id = request.GET.get("turma", "")

    inscricoes = Inscricao.objects.filter(
        situacao=Inscricao.Situacao.ATIVA
    ).select_related(
        "cursista",
        "cursista__escola",
        "turma",
        "turma__formacao",
    ).order_by(
        "cursista__nome_completo"
    )

    if formacao_id.isdigit():
        inscricoes = inscricoes.filter(
            formacao_id=int(formacao_id)
        )

    if turma_id.isdigit():
        inscricoes = inscricoes.filter(
            turma_id=int(turma_id)
        )

    resposta = HttpResponse(content_type="application/pdf")
    resposta["Content-Disposition"] = (
        'attachment; filename="relatorio_participantes.pdf'
    )
    documento = SimpleDocTemplate(
        resposta,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    estilos = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle(
        "TituloRelatorio",
        parent=estilos["title"],
        alignment=TA_CENTER,
        fontSize=16,
        leading=20,
        spaceAfter=8,
    )

    texto_estilo = ParagraphStyle(
        "TextoRelatorio",
        parent=estilos["Normal"],
        fontSize=8,
        leading=10,
    )

    celula_estilo = ParagraphStyle(
        "CelulaTabela",
        parent=estilos["Normal"],
        fontSize=7,
        leading=9,
    )

    conteudo = []

    conteudo.append(
        Paragraph(
            "Relatório de Participantes das Formações Continuadas",
            titulo_estilo,
        )
    )
    conteudo.append(
        Paragraph(
            f"Emitido em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}",
            texto_estilo,
        )
    )

    formacao_selecionada = None
    turma_selecionada = None

    if formacao_id.isdigit():
        formacao_selecionada = Formacao.objects.filter(
            pk=int(formacao_id)
        ).first()

    if turma_id.isdigit():
        turma_selecionada = Turma.objects.select_related(
            "formacao"
        ).filter(pk=int(turma_id)).first()

    nome_formacao = (
        formacao_selecionada.titulo if formacao_selecionada else "Todas as formações"
    )

    nome_turma = (
        turma_selecionada.nome if turma_selecionada else "Todas as turmas"
    )

    conteudo.append(
        Paragraph(
            f"Formação: {escape(nome_formacao)}", texto_estilo,
        )
    )
    conteudo.append(
        Paragraph(
            f"Turma: {escape(nome_turma)}", texto_estilo,
        )
    )
    conteudo.append(Spacer(1,6*mm))

    cabecalhos = [
        "Participante",
        "CPF",
        "Escola",
        "Formação",
        "Turma",
        "Horas frequentadas",
        "Frequência",
        "Resultado",
    ]
    dados = [[
        Paragraph(cabecalho, celula_estilo)
        for cabecalho in cabecalhos
    ]]

    total_participantes = 0
    total_concluintes = 0

    for inscricao in inscricoes:
        print(inscricao)
        calculo = calcular_frequencia(inscricao)
        cursista = inscricao.cursista

        if calculo["horas_realizadas"] == 0:
            resultado = "Sem encontros realizados"
        elif calculo["concluinte"]:
            resultado = "Concluinte"
            total_concluintes += 1
        else:
            resultado = "Não concluinte"

        total_participantes += 1
        dados.append([
            Paragraph(
                escape(cursista.nome_completo),
                celula_estilo,
            ),
            Paragraph(escape(cursista.cpf), celula_estilo),
            Paragraph(escape(cursista.escola.nome), celula_estilo),
            Paragraph(escape(inscricao.turma.formacao.titulo), celula_estilo),
            Paragraph(escape(inscricao.turma.nome), celula_estilo),
            str(calculo["horas_realizadas"]),
            f'{calculo["percentual_frequencia"]:.2f}%',
            Paragraph(escape(resultado), celula_estilo),
        ])

    if total_participantes == 0:
        dados.append([
            Paragraph("Nenhum participante encontrado.", celula_estilo),
            "", "", "", "", "", "", "",
        ])

    conteudo.append(
        Paragraph(
            f"Total de participantes: {total_participantes} | Concluintes: {total_concluintes}",
            texto_estilo
        )
    )

    conteudo.append(Spacer(1,4 *mm))
    tabela = Table(
        dados,
        colWidths=[
            40 * mm,
            25 * mm,
            40 * mm,
            45 * mm,
            30 * mm,
            25 * mm,
            22 * mm,
            35 * mm,
        ], repeatRows=1,
    )

    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#169034")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor("#F2F6F3"),
        ]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))

    conteudo.append(tabela)
    documento.build(conteudo)

    return resposta

@staff_member_required
def relatorio_consolidado(request):
    formacoes = Formacao.objects.all().order_by("titulo")
    resultados = []
    total_formacoes = formacoes.count()
    total_horas_gerais = 0
    total_concluintes_geral = 0

    for formacao in formacoes:
        total_turmas = formacao.turmas.count()
        encontros_realizados = Encontro.objects.filter(
            turma__formacao=formacao,
            situacao=Encontro.Situacao.REALIZADO,
        )
        carga_horaria = (
            encontros_realizados.aggregate(
                total=Sum("carga_horaria")
            )["total"] or 0
        )

        inscricoes = Inscricao.objects.filter(
            turma__formacao = formacao,
            situacao=Inscricao.Situacao.ATIVA
        ).select_related(
            "cursista",
            "turma",
            "turma__formacao",
        )
        total_inscricoes = inscricoes.count()
        total_concluintes = 0

        for inscricao in inscricoes:
            calculo = calcular_frequencia(inscricao)
            if calculo['concluinte']:
                total_concluintes += 1

        resultados.append({
            "formacao": formacao,
            "total_turmas": total_turmas,
            "total_encontros": encontros_realizados.count(),
            "carga_horaria": carga_horaria,
            "total_inscricoes": total_inscricoes,
            "total_concluintes": total_concluintes,
        })

        total_horas_gerais += carga_horaria
        total_concluintes_geral += total_concluintes

    contexto = {
        "resultados": resultados,
        "total_formacoes": total_formacoes,
        "total_horas_gerais": total_horas_gerais,
        "total_concluintes_geral": total_concluintes_geral,
    }

    return render(
        request,
        "relatorios/consolidado.html",
        contexto,
    )

@staff_member_required
def export_consolidado_excel(request):
    formacoes = Formacao.objects.all().order_by("titulo")

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Relatório Consolidado"

    planilha.append([
        "Formação",
        "Turmas",
        "Encontros realizados",
        "Horas ministradas",
        "Inscrições ativas",
        "Concluintes",
    ])

    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor="169034",
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    total_formacoes = 0
    total_horas = 0
    total_concluintes = 0

    for formacao in formacoes:
        encontros_realizados = Encontro.objects.filter(
            turma__formacao=formacao,
            situacao=Encontro.Situacao.REALIZADO,
        )

        carga_horaria = (
            encontros_realizados.aggregate(
                total=Sum("carga_horaria")
            )["total"] or 0
        )

        inscricoes = Inscricao.objects.filter(
            turma__formacao=formacao,
            situacao=Inscricao.Situacao.ATIVA,
        ).select_related(
            "cursista",
            "turma",
        )

        concluintes = sum(
            1 for inscricao in inscricoes if calcular_frequencia(inscricao)["concluinte"]
        )
        planilha.append([
            formacao.titulo,
            formacao.turmas.count(),
            encontros_realizados.count(),
            float(carga_horaria),
            inscricoes.count(),
            concluintes,
        ])
        total_formacoes += 1
        total_horas += carga_horaria
        total_concluintes += concluintes

    planilha.append([])
    planilha.append([
        "TOTAIS",
        "",
        "",
        float(total_horas),
        "",
        total_concluintes,
    ])

    for celula in planilha[planilha.max_row]:
        celula.font = Font(bold=True)
    larguras = {
        "A": 42,
        "B": 14,
        "C": 24,
        "D": 22,
        "E": 20,
        "F": 16,
    }
    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = f"A1:f{planilha.max_row - 2}"

    resposta = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
    resposta["Content-Disposition"] = (
        'attachment; filename="relatorio_consolidado.xlsx"'
    )

    workbook.save(resposta)
    return resposta

@staff_member_required
def export_consolidado_pdf(request):
    formacoes = Formacao.objects.all().order_by("titulo")
    resposta = HttpResponse(content_type="application/pdf")
    resposta["Content-Disposition"] = ('attachment; filename="relatorio_consolidado.pdf"')

    documento = SimpleDocTemplate(
        resposta,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )
    estilos = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle(
        "TituloConsolidado",
        parent=estilos["Title"],
        alignment=TA_CENTER,
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#169034"),
        spaceAfter=8,
    )
    texto_estilo = ParagraphStyle(
        "TextoConsolidado",
        parent=estilos["Normal"],
        fontSize=9,
        leading=12,
    )
    celula_estilo = ParagraphStyle(
        "CelulaConsolidado",
        parent=estilos["Normal"],
        fontSize=8,
        leading=10,
    )

    conteudo = [
        Paragraph("Relatório Consolidado das Formações Continuadas", titulo_estilo),
        Paragraph("Secretaria Municipal de Educação", texto_estilo),
        Paragraph(f"Emitido em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}", texto_estilo),
        Spacer(1, 8 * mm),
    ]

    cabecalhos = [
        "Formação",
        "Turmas",
        "Encontros realizados",
        "Horas ministradas",
        "Inscrições ativas",
        "Concluintes",
    ]

    dados = [[
        Paragraph(cabecalho, celula_estilo) for cabecalho in cabecalhos
    ]]

    total_formacoes = 0
    total_horas = 0
    total_inscricoes = 0
    total_concluintes = 0

    for formacao in formacoes:
        encontros = Encontro.objects.filter(
            turma__formacao=formacao,
            situacao=Encontro.Situacao.REALIZADO
        )
        carga_horaria = (
            encontros.aggregate(
                total=Sum("carga_horaria")
            )["total"] or 0
        )
        inscricoes = Inscricao.objects.filter(
            turma__formacao=formacao,
            situacao=Inscricao.Situacao.ATIVA,
        ).select_related(
            "cursista",
            "turma",
        )

        concluintes = sum(
            1 for inscricao in inscricoes if calcular_frequencia(inscricao)["concluinte"]
        )

        dados.append([
            Paragraph(formacao.titulo, celula_estilo),
            str(formacao.turmas.count()),
            str(encontros.count()),
            f"{carga_horaria:.2f}",
            str(inscricoes.count()),
            str(concluintes),
        ])

        total_formacoes += 1
        total_horas += carga_horaria
        total_inscricoes += inscricoes.count()
        total_concluintes += concluintes

    if total_formacoes == 0:
        dados.append([
            Paragraph("Nenhuma formação cadastrada.", celula_estilo),
            "", "", "", "", "",
        ])

    dados.append([
        Paragraph("Totais", celula_estilo),
        "",
        "",
        f"{total_horas:.2f}",
        str(total_inscricoes),
        str(total_concluintes),
    ])
    conteudo.append(
        Paragraph(
            f"Formações cadastradas: {total_formacoes} "
            f"| Horas ministradas: {total_horas:.2f} "
            f"| Inscrições concluídas: {total_concluintes}",
            texto_estilo,
        )
    )
    conteudo.append(Spacer(1,4*mm))

    tabela = Table(
        dados,
        colWidths=[
            75 * mm,
            20 * mm,
            35 * mm,
            30 * mm,
            30 * mm,
            25 * mm,
        ],
        repeatRows=1,
    )

    tabela.setStyle(TableStyle([
        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#169034"),
        ),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#F2F6F3")],
        ),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    conteudo.append(tabela)
    documento.build(conteudo)

    return resposta